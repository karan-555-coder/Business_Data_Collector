"""Unit tests for the performance changes (no network, no credits):
shared TLS context, change-only Excel exports, per-category stats cache.
Run:  python tests/test_performance.py  (from the app/ directory)"""

from __future__ import annotations

import os
import ssl
import sys
import tempfile
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import requests  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from backend.collector import exporter  # noqa: E402
from backend.collector.crawler import make_session  # noqa: E402
from backend.collector.engine import StateStore  # noqa: E402
from backend.collector.search import _serper_session  # noqa: E402
from backend.collector.tls import SharedTLSAdapter, shared_tls_context  # noqa: E402


def rec(i: int) -> dict:
    return {"Company Name": f"Firm {i}", "Official Website": f"https://firm{i}.com",
            "Business Email": f"info@firm{i}.com" if i % 2 else "",
            "Business Phone": f"+1 555 {i:04d}", "City": "Austin"}


def test_shared_tls():
    ctx = shared_tls_context()
    assert ctx is shared_tls_context(), "context must be built once"
    assert ctx.verify_mode == ssl.CERT_REQUIRED
    assert ctx.cert_store_stats()["x509_ca"] > 0, "CA bundle not loaded"
    for s in (make_session(), _serper_session()):
        ad = s.get_adapter("https://example.com/")
        assert isinstance(ad, SharedTLSAdapter)
        req = requests.Request("GET", "https://example.com/").prepare()
        _, kw = ad.build_connection_pool_key_attributes(req, True, None)
        assert kw.get("ssl_context") is ctx, "verified HTTPS pools must share the context"
        _, kw = ad.build_connection_pool_key_attributes(req, False, None)
        assert "ssl_context" not in kw, "verify=False must keep requests' own handling"
        _, kw = ad.build_connection_pool_key_attributes(req, "/custom/ca.pem", None)
        assert "ssl_context" not in kw, "a custom CA bundle must keep requests' own handling"
        http = requests.Request("GET", "http://example.com/").prepare()
        _, kw = ad.build_connection_pool_key_attributes(http, True, None)
        assert "ssl_context" not in kw

        class Conn:
            cert_reqs = ca_certs = ca_cert_dir = None
        c = Conn()
        ad.cert_verify(c, "https://example.com/", True, None)
        assert c.cert_reqs == "CERT_REQUIRED" and c.ca_certs is None, \
            "the bundle path must not be re-loaded per connection"
    print("  shared TLS context: one verified context for every HTTPS pool")


def test_excel_only_when_changed():
    d = tempfile.mkdtemp(prefix="perf_test_")
    path = os.path.join(d, "01_Test.xlsx")
    recs = [rec(i) for i in range(50)]
    exporter.write_category_file("Test", "01_Test.xlsx", recs, d)
    st1 = os.stat(path)
    time.sleep(0.05)
    exporter.write_category_file("Test", "01_Test.xlsx", list(recs), d)
    assert os.stat(path).st_mtime_ns == st1.st_mtime_ns, "unchanged rows must not be rewritten"
    exporter.write_category_file("Test", "01_Test.xlsx", recs, d, force=True)
    assert os.stat(path).st_mtime_ns != st1.st_mtime_ns, "force must rewrite"
    st2 = os.stat(path)
    recs.append(rec(50))
    exporter.write_category_file("Test", "01_Test.xlsx", recs, d)
    assert os.stat(path).st_mtime_ns != st2.st_mtime_ns, "new rows must be written"
    ws = load_workbook(path).active
    assert ws.max_row == 52 and ws.cell(row=52, column=1).value == "Firm 50"
    os.remove(path)                                  # e.g. deleted by hand
    exporter.write_category_file("Test", "01_Test.xlsx", recs, d)
    assert os.path.exists(path), "a removed file must be regenerated"
    exporter.write_category_file("Renamed", "01_Test.xlsx", recs, d)
    assert load_workbook(path).active.title == "Renamed", "sheet title is content"

    cats = {"Test": {"display": "Test", "file": "01_Test.xlsx"}}
    stats = {"Test": {"discovered": 60, "duplicates": 3, "failed_urls": 2}}
    spath = os.path.join(d, exporter.SUMMARY_FILE)
    exporter.write_master_summary(cats, {"Test": recs}, stats, d)
    s1 = os.stat(spath)
    time.sleep(0.05)
    exporter.write_master_summary(cats, {"Test": recs}, dict(stats), d)
    assert os.stat(spath).st_mtime_ns == s1.st_mtime_ns, "unchanged summary rewritten"
    stats["Test"]["failed_urls"] = 3
    exporter.write_master_summary(cats, {"Test": recs}, stats, d)
    ws = load_workbook(spath).active
    assert ws.cell(row=2, column=9).value == 3 and ws.cell(row=3, column=1).value == "TOTAL"
    print("  Excel: rewritten only when content changed (force, delete, title handled)")


def test_category_stats_cache():
    st = StateStore()
    cat = next(iter(st.records))
    other = list(st.records)[1]
    assert st.category_stats()[cat]["count"] == 0
    for i in range(5):
        assert st.append_record(cat, rec(i), 100)
    s = st.category_stats()
    assert s[cat] == {"count": 5, "emails": 2, "websites": 5}, s[cat]
    st.append_record(other, rec(7), 100)
    s = st.category_stats()
    assert s[other]["count"] == 1 and s[cat]["count"] == 5
    st.delete_category(cat)
    assert st.category_stats()[cat]["count"] == 0, "delete must invalidate the cache"
    print("  category stats: per-category cache stays exact")


def test_timeouts_from_env():
    import subprocess
    code = ("from backend import config as c; print(c.FETCH_TIMEOUT, c.ROBOTS_TIMEOUT, "
            "c.SERPER_TIMEOUT, c.FETCH_DEADLINE)")
    app_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    base = {k: v for k, v in os.environ.items()
            if not k.endswith(("_TIMEOUT", "FETCH_DEADLINE"))}
    out = subprocess.run([sys.executable, "-c", code], cwd=app_dir, env=base,
                         capture_output=True, text=True, check=True).stdout.strip()
    assert out == "(5.0, 12.0) (4.0, 8.0) (5.0, 30.0) 25.0", f"defaults changed: {out}"
    env = dict(base, FETCH_CONNECT_TIMEOUT="3", FETCH_READ_TIMEOUT="9",
               ROBOTS_CONNECT_TIMEOUT="2", ROBOTS_READ_TIMEOUT="5",
               SERPER_CONNECT_TIMEOUT="4", SERPER_READ_TIMEOUT="20",
               FETCH_DEADLINE="999")                       # clamped to 300
    out = subprocess.run([sys.executable, "-c", code], cwd=app_dir, env=env,
                         capture_output=True, text=True, check=True).stdout.strip()
    assert out == "(3.0, 9.0) (2.0, 5.0) (4.0, 20.0) 300.0", out
    print("  timeouts: env-configurable, defaults unchanged, values clamped")


def test_health_alias():
    import asyncio
    from backend import main as m
    assert m.FAST_ROUTES["/health"] is m.FAST_ROUTES["/api/health"]
    r = asyncio.run(m.FAST_ROUTES["/health"]({}, {}, "t"))
    assert r.status == 200 and b'"status":"ok"' in r.body.replace(b" ", b"")
    print("  /health: same lightweight probe as /api/health")


def test_priority_crawl_pool():
    import threading
    from backend.collector.engine import PRIO_MINED, PRIO_PLACE, PRIO_RESULT
    from backend.collector.prio_pool import PriorityThreadPool
    pool = PriorityThreadPool(max_workers=1)
    gate, order = threading.Event(), []
    pool.submit(gate.wait)                     # occupy the only worker
    futs = [pool.submit(order.append, name, prio=p) for name, p in (
        ("mined1", PRIO_MINED), ("result1", PRIO_RESULT), ("place1", PRIO_PLACE),
        ("mined2", PRIO_MINED), ("place2", PRIO_PLACE))]
    gate.set()
    for f in futs:
        f.result(timeout=5)
    assert order == ["place1", "place2", "result1", "mined1", "mined2"], order
    blocker = threading.Event()
    pool.submit(blocker.wait)
    queued = pool.submit(order.append, "never", prio=PRIO_PLACE)
    pool.shutdown(wait=False, cancel_futures=True)   # what a finished job does
    blocker.set()
    assert queued.cancelled() and "never" not in order
    print("  crawl pool: places -> search results -> mined links, FIFO within, "
          "cancel on shutdown")


def test_analysis_lane():
    import threading
    from backend import config
    from backend.collector import analysis
    if config.ANALYZE_PROCESSES > 0:
        print("  analysis lane: skipped (process pool in use on this machine)")
        return
    busy, peak, lock = [0], [0], threading.Lock()
    real = analysis.analyze_html

    def probe(*a):
        with lock:
            busy[0] += 1
            peak[0] = max(peak[0], busy[0])
        time.sleep(0.02)
        with lock:
            busy[0] -= 1
        return {}
    analysis.analyze_html = probe
    try:
        ts = [threading.Thread(target=analysis.analyze, args=(b"<p>x</p>", "https://a.com"))
              for _ in range(12)]
        [t.start() for t in ts]
        [t.join() for t in ts]
    finally:
        analysis.analyze_html = real
    assert peak[0] <= config.ANALYZE_THREADS, (peak[0], config.ANALYZE_THREADS)
    print(f"  analysis lane: at most {config.ANALYZE_THREADS} page(s) analysed at once")


def main():
    test_shared_tls()
    test_excel_only_when_changed()
    test_category_stats_cache()
    test_timeouts_from_env()
    test_health_alias()
    test_priority_crawl_pool()
    test_analysis_lane()
    print("\nALL PERFORMANCE TESTS PASSED")


if __name__ == "__main__":
    main()

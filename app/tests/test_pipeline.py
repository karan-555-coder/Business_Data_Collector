"""End-to-end pipeline test WITHOUT Serper credits or network access.

A fake search provider and fake fetcher feed canned HTML through the real
engine: crawl -> extract -> validate -> dedup -> categorize -> export.
Run:  python tests/test_pipeline.py   (from the app/ directory)
"""

from __future__ import annotations

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend import config  # noqa: E402

_tmp = tempfile.mkdtemp(prefix="collector_test_")
config.OUTPUT_DIR = _tmp
config.STATE_PATH = os.path.join(_tmp, "demo_state.json")

from backend.collector import discovery, engine  # noqa: E402
from backend.collector.engine import CollectionJob, StateStore  # noqa: E402
from backend.collector.search import SearchProvider  # noqa: E402
from bs4 import BeautifulSoup  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

engine.LIST_PAGE_MIN_EXTERNAL_LINKS = 1  # let the tiny fixture count as a list page

PAGES = {
    "http://acmetalent.example": """
      <html><head><title>Acme Talent Solutions | RPO Services</title>
      <meta name="description" content="End-to-end recruitment process outsourcing."></head>
      <body><a href="mailto:hello@acmetalent.example">Email</a>
      <a href="tel:+912233445566">Call</a>
      <address>45 Business Park, Mumbai, Maharashtra 400001, India</address>
      <a href="/services/rpo">RPO Sourcing</a></body></html>""",
    "http://bizjournal.example": """
      <html><head><title>Top 10 RPO Companies 2026 | Business Journal Times</title></head>
      <body><a href="mailto:editor@bizjournal.example">e</a>
      <a href="tel:+911112223334">t</a><p>Ranking article, not a company.</p></body></html>""",
    "http://listhub.example/best-rpo": """
      <html><head><title>Best RPO Firms</title></head>
      <body><a href="http://nimbushiring.example">Nimbus Hiring Solutions</a></body></html>""",
    "http://nimbushiring.example": """
      <html><head><title>Nimbus Hiring Solutions - Recruitment Outsourcing</title></head>
      <body><a href="mailto:contact@nimbushiring.example">mail</a>
      <a href="tel:+914455667788">call</a></body></html>""",
    "http://vertexrpo.example": """
      <html><head><title>Vertex RPO Services</title></head>
      <body><a href="mailto:info@vertexrpo.example">mail</a></body></html>""",
    "http://shieldnet.example": """
      <html><head><title>ShieldNet Security | Managed Cybersecurity</title></head>
      <body><a href="mailto:soc@shieldnet.example">mail</a>
      <a href="tel:+13105551234">call</a>
      <address>800 Wilshire Blvd, Los Angeles, California 90017, USA</address></body></html>""",
    "http://pixelforge.example": """
      <html><head><title>PixelForge Studio | 3D Visualization</title></head>
      <body><a href="mailto:studio@pixelforge.example">mail</a>
      <a href="tel:+918899001122">call</a>
      <address>7 Design Tower, Pune, Maharashtra 411001, India</address></body></html>""",
}


class FakeFetcher:
    def fetch_raw(self, url: str):
        """What the engine calls now (parsing happens in analysis.analyze)."""
        html = PAGES.get(url.rstrip("/"))
        if html is None:
            return None, url, "connection/DNS error"
        return html.encode("utf-8"), url, ""

    def fetch_page(self, url: str):
        html = PAGES.get(url.rstrip("/"))
        if html is None:
            return None, "", url, "connection/DNS error"
        return BeautifulSoup(html, "html.parser"), html, url, ""


class FakeProvider(SearchProvider):
    name = "fake"
    supports_places = True
    credits_used = 0

    def __init__(self, organic_urls, places_rows):
        self._organic = organic_urls
        self._places = places_rows
        self._organic_served = False
        self._places_served = False
        self.credits_used = 0     # 1 credit per call, like Serper

    def organic(self, query, max_results):
        self.credits_used += 1
        if self._organic_served:
            return []
        self._organic_served = True
        return list(self._organic)

    def places(self, query, pages):
        self.credits_used += 1
        if self._places_served:
            return []
        self._places_served = True
        return list(self._places)


def planner_tests():
    """Discovery planner: no trivial rephrasings, geographic spread, Places
    follow-up pages, persistence across runs."""
    from backend.collector.categories import search_phrases
    from backend.collector.discovery import DiscoveryPlanner, canon_phrase
    from backend.collector.geo import expansion_geos, find_state

    assert canon_phrase("3D rendering studios") == canon_phrase("3D rendering services")
    assert canon_phrase("top 3D studios") == canon_phrase("3D studios")
    assert canon_phrase("best management consulting firms") == \
        canon_phrase("management consulting companies")
    assert canon_phrase("3D animation studio") != canon_phrase("3D rendering studio")

    geos = expansion_geos("", "Texas", "USA", "Texas, USA", [])
    assert geos[0] == "Texas, USA" and "Houston, Texas" in geos and len(geos) > 10
    assert find_state("TX", "USA")[1] == "Texas"
    assert expansion_geos("Austin", "Texas", "USA", "Austin, Texas, USA", []) == \
        ["Austin, Texas, USA"], "pinned city must not expand"
    g_us = expansion_geos("", "", "USA", "USA", ["New York", "Chicago"])
    assert g_us[:3] == ["USA", "New York", "Chicago"] and len(g_us) > 100
    assert not any(g.startswith("Chicago,") for g in g_us), "major city repeated"

    phrases = search_phrases("3D_Studios", "3D Studios", ["architectural rendering"])
    pl = DiscoveryPlanner(phrases, geos, places=True, max_page=3)
    assert pl.dropped_rephrasings >= 1          # "3D rendering services"
    issued = []
    for _ in range(8):
        reqs = pl.next_requests(want_organic=True)
        assert {r.kind for r in reqs} == {"places", "organic"}
        issued.append(reqs[0])
    # breadth first: 8 searches land in 8 different locations
    assert len({r.geo for r in issued}) == 8, [r.query for r in issued]
    assert issued[0].query == "architectural rendering in Texas, USA"
    texts = [r.query.lower() for r in issued]
    assert len(set(texts)) == len(texts)
    # a mostly-new Places page earns a page-2 follow-up, valued at the
    # LEARNED page-2 novelty (live: 1/3 of page 1), so untried cells go first
    # (once the category's few exploratory follow-ups have been measured)
    pl.fu = [0.0, 0, discovery.FOLLOWUP_EXPLORE]
    pl.on_result(issued[0], results=10, new=9)
    nxt = pl.next_requests(want_organic=True)
    assert nxt[0].page == 1, "page 2 must not outrank an untried cell by default"
    assert len(pl.followups) == 1 and pl.followups[0].page == 2
    # ... but once page 2 is measured to be as novel as page 1, it wins
    pl.fu = [30.0, 45, 3]
    pl.followups[0].expected = 0.9 * pl._followup_ratio()
    nxt = pl.next_requests(want_organic=True)
    assert nxt[0].page == 2 and nxt[0].query == issued[0].query
    # a thin Places result (few listings) of a Places-only cell queues an
    # organic search of that cell; a cell already searched organically doesn't
    pl.on_result(issued[1], results=2, new=2)
    assert not any(r.kind == "organic" for r in pl.followups)
    solo = pl.next_requests(want_organic=False)
    assert [r.kind for r in solo] == ["places"]
    pl.on_result(solo[0], results=2, new=2)
    assert any(r.kind == "organic" and r.query == solo[0].query
               for r in pl.followups)
    # saturated locations score below untried ones
    for r in issued[1:]:
        pl.on_result(r, results=10, new=0)
    nxt = pl.next_requests(want_organic=True)
    assert nxt[0].geo not in {r.geo for r in issued}
    # persistence: a resumed planner never re-issues a used cell
    pl2 = DiscoveryPlanner(phrases, geos, saved=pl.export(), places=True)
    again = set()
    while True:
        reqs = pl2.next_requests(want_organic=False)
        if not reqs:
            break
        again.add(reqs[0].query.lower())
        pl2.on_result(reqs[0], results=10, new=5)
    assert not again & set(texts), "used cell re-issued after resume"
    assert len(again) > 50, "planner ran out of new searches too early"
    # queries executed by older versions (done_queries) are honoured
    pl3 = DiscoveryPlanner(phrases, geos, legacy_done={texts[0]}, places=True)
    assert pl3.next_requests(want_organic=False)[0].query.lower() != texts[0]
    print("planner tests passed")


def run_job(state, category, keywords, organic, places, location="India", geo=None):
    job = CollectionJob(state, category, keywords, location, 20, "serper", 4, geo=geo)
    job.fetcher = FakeFetcher()
    real_make = engine.make_provider
    engine.make_provider = lambda *_: FakeProvider(organic, places)
    try:
        job._run()
    finally:
        engine.make_provider = real_make
    return job


def main():
    planner_tests()
    # Query pool expansion: base queries, then cities in the country, then
    # modifier variants — enough material to keep searching to the target.
    from backend.collector.categories import CATEGORIES, build_query_pool
    tpls = CATEGORIES["Finance"]["templates"]
    pool = build_query_pool(tpls, ["CPA"], "India", "India", True, 40)
    assert len(pool) == 40, f"pool not filled: {len(pool)}"
    assert any("Mumbai" in q for q in pool), "city expansion missing"
    assert len(set(q.lower() for q in pool)) == len(pool), "duplicate queries"
    # without city expansion, modifier variants fill the pool instead
    pool_nc = build_query_pool(tpls, [], "Portugal", "Portugal", True, 20)
    assert any(q.startswith(("top ", "best ")) for q in pool_nc), \
        "modifier expansion missing"
    # deep pool: a 2000-record target must not run out of query variations
    cds = CATEGORIES["CDS_Corporate_Compliance"]["templates"]
    deep = build_query_pool(cds, ["company secretary"], "India", "India", True, 400)
    assert len(deep) >= 300, f"deep pool too small: {len(deep)}"
    assert len(set(q.lower() for q in deep)) == len(deep)
    # a pinned city must NOT expand to other cities
    pool2 = build_query_pool(tpls, [], "Los Angeles, California, USA", "USA", False, 40)
    assert not any("Chicago" in q for q in pool2), "city expansion leaked scope"

    state = StateStore()

    rpo_places = [
        {"title": "Vertex RPO Services", "website": "http://vertexrpo.example",
         "phoneNumber": "+91 98765 43210", "cid": "111",
         "address": "12 MG Road, Bangalore, Karnataka 560001, India",
         "category": "Recruitment agency"},
        # duplicate of the row above by domain and name
        {"title": "Vertex RPO Services", "website": "http://vertexrpo.example",
         "phoneNumber": "+91 98765 43210", "cid": "112",
         "address": "12 MG Road, Bangalore, Karnataka 560001, India",
         "category": "Recruitment agency"},
        # no website: valid Places-only record keyed on cid
        {"title": "Summit Staffing Bureau", "phoneNumber": "+91 91234 56780",
         "cid": "113", "address": "3 Anna Salai, Chennai, Tamil Nadu 600002, India",
         "category": "Recruitment agency"},
    ]
    rpo_organic = [
        "http://acmetalent.example",       # valid company
        "http://bizjournal.example",       # publisher -> must be rejected
        "http://listhub.example/best-rpo", # directory -> mine Nimbus from it
        "http://dead-site.example",        # fetch failure -> failed URL
    ]
    job = run_job(state, "RPO", ["recruitment process outsourcing"], rpo_organic, rpo_places)
    rpo = state.records["RPO"]
    names = {r["Company Name"] for r in rpo}
    print("RPO records:", sorted(names))
    assert job.status == "exhausted", job.error
    assert "Acme Talent Solutions" in names
    assert "Nimbus Hiring Solutions" in names, "directory mining failed"
    assert "Vertex RPO Services" in names
    assert "Summit Staffing Bureau" in names, "places-only record missing"
    assert not any("Journal" in n for n in names), "publisher not rejected"
    assert job.counters["duplicates"] >= 1, "duplicate not detected"
    assert job.counters["failed"] >= 1, "failed URL not counted"
    acme = next(r for r in rpo if r["Company Name"] == "Acme Talent Solutions")
    assert acme["Business Email"] == "hello@acmetalent.example"
    assert acme["City"] == "Mumbai" and acme["State"] == "Maharashtra" and acme["Country"] == "India"
    assert acme["Category"] == "RPO" and int(acme["Confidence Score"]) >= 70
    vertex = next(r for r in rpo if r["Company Name"] == "Vertex RPO Services")
    assert vertex["Business Email"] == "info@vertexrpo.example", "place website not crawled for email"

    # Second category: 3D Studios must stay in its own file.
    job2 = run_job(state, "3D_Studios", ["3D visualization"],
                   ["http://pixelforge.example"], [])
    assert job2.status == "exhausted", job2.error
    threed = {r["Company Name"] for r in state.records["3D_Studios"]}
    print("3D records:", sorted(threed))
    assert "PixelForge Studio" in threed

    # Cross-category dedup: Acme rediscovered under 3D must be rejected.
    job3 = run_job(state, "3D_Studios", [], ["http://acmetalent.example"], [])
    assert job3.status == "exhausted", job3.error
    assert "Acme Talent Solutions" not in {r["Company Name"] for r in state.records["3D_Studios"]}

    # Excel outputs: separate files, no category mixing, summary present.
    rpo_wb = load_workbook(os.path.join(config.OUTPUT_DIR, "07_RPO.xlsx"))
    rows = list(rpo_wb.active.iter_rows(min_row=2, values_only=True))
    header = [c.value for c in rpo_wb.active[1]]
    assert header[0] == "Company Name" and "Confidence Score" in header
    assert len(rows) == len(rpo) and all(r[1] == "RPO" for r in rows)
    from backend.collector.categories import CATEGORIES
    threed_wb = load_workbook(os.path.join(config.OUTPUT_DIR, "09_3D_Studios.xlsx"))
    trows = list(threed_wb.active.iter_rows(min_row=2, values_only=True))
    threed_display = CATEGORIES["3D_Studios"]["display"]
    assert all(r[1] == threed_display for r in trows) and len(trows) == len(threed)
    assert os.path.exists(os.path.join(config.OUTPUT_DIR, "00_Master_Summary.xlsx"))

    # Custom category with structured location (Country + State + City).
    slug = state.ensure_custom("Cybersecurity companies")
    jobc = run_job(state, slug, ["network security"],
                   ["http://shieldnet.example"], [],
                   location="Los Angeles, California, USA",
                   geo={"city": "Los Angeles", "state": "California", "country": "USA"})
    assert jobc.status == "exhausted", jobc.error
    cyber = state.records[slug]
    assert len(cyber) == 1 and cyber[0]["Company Name"] == "ShieldNet Security"
    assert cyber[0]["Category"] == "Cybersecurity companies"
    assert cyber[0]["City"] == "Los Angeles" and cyber[0]["State"] == "California"
    assert cyber[0]["Country"] == "USA"
    assert int(cyber[0]["Confidence Score"]) >= 70, "custom terms not matched"
    cfile = state.all_categories()[slug]["file"]
    cyber_wb = load_workbook(os.path.join(config.OUTPUT_DIR, cfile))
    crows = list(cyber_wb.active.iter_rows(min_row=2, values_only=True))
    assert len(crows) == 1 and crows[0][1] == "Cybersecurity companies"

    # Checkpoint/resume: reload state from disk and confirm nothing is lost.
    reloaded = StateStore.load(config.STATE_PATH)
    assert {r["Company Name"] for r in reloaded.records["RPO"]} == names
    assert len(reloaded.done_queries) >= 2
    assert slug in reloaded.custom and len(reloaded.records[slug]) == 1, \
        "custom category lost on reload"
    added, dup_of = reloaded.registry.check_and_add(
        {"Company Name": "Acme Talent Solutions",
         "Official Website": "http://acmetalent.example"}, "Finance")
    assert not added, "reloaded registry lost dedup keys"

    print("\nALL PIPELINE TESTS PASSED")
    print("outputs in:", config.OUTPUT_DIR)


if __name__ == "__main__":
    main()
